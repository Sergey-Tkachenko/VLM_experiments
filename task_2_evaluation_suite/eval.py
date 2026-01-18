from __future__ import annotations

import argparse
import json
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import torch
from dotenv import load_dotenv
from loguru import logger
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from tqdm import tqdm

REPO_ROOT = Path(__file__).resolve().parents[1]

load_dotenv(REPO_ROOT / ".env")

sys.path.insert(0, str(REPO_ROOT))

from task_2_evaluation_suite.config_loader import EvalConfig, load_eval_config
from task_1_sample_inference.inference import InferenceOutcome, QwenVideoInferencer
from task_1_sample_inference.prompting import load_prompt_template, render_prompt




PROMPT_PATH = Path(__file__).resolve().parent / "prompt.txt"
CLASS_MAPPING_PATH = Path(__file__).resolve().parent / "type_id_to_parsed_description_mapping.json"
POPULAR_CLASS_IDS = [11, 43, 50, 10, 5, 6, 37, 48, 38, 8, 1, 57, 12, 49, 56, 14, 39, 42, 9]


class DadaAccidentSchema(BaseModel):
    """Schema for DADA accident predictions."""

    accident_type: int = Field(
        ...,
        description="Accident type ID. Use 0 if no accident occurs.",
        ge=0,
    )
    accident_frame_position_sec: float | None = Field(
        None,
        description="Accident frame timestamp in seconds from clip start. Null if no accident occurs.",
    )


@dataclass(frozen=True)
class ClipMeta:
    """Metadata for a single clip."""

    type_id: str
    video_id: str
    clip_id: str
    accident_frame: int | None


def _normalize_header(header: str) -> str:
    return " ".join(str(header).strip().lower().split())


def _normalize_type_id(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    try:
        return str(int(value))
    except (TypeError, ValueError):
        return None


def _normalize_video_id(value: Any, target_width: int = 3) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if text.isdigit():
        return text.zfill(target_width)
    return text


def _parse_split_item(item: list[Any]) -> tuple[str, str, str]:
    if not isinstance(item, list) or len(item) < 2:
        raise ValueError("Invalid split item structure.")
    pair = item[0]
    clip_id = str(item[1]).strip()
    if not isinstance(pair, list) or len(pair) != 2:
        raise ValueError("Split item must contain [type_id, video_id].")
    type_id = str(pair[0]).strip()
    video_id = _normalize_video_id(pair[1])
    return type_id, video_id, clip_id


def load_split(split_path: Path) -> list[ClipMeta]:
    """Load split metadata entries from JSON.

    Args:
        split_path: Path to a split JSON file.

    Returns:
        List of clip metadata entries.
    """
    if not split_path.exists():
        raise FileNotFoundError(f"Split not found: {split_path}")
    raw = json.loads(split_path.read_text(encoding="utf-8"))
    clips = []
    for item in raw:
        type_id, video_id, clip_id = _parse_split_item(item)
        clips.append(ClipMeta(type_id=type_id, video_id=video_id, clip_id=clip_id, accident_frame=None))
    return clips


def _find_header_index(headers: list[str], needle: str) -> int | None:
    for idx, header in enumerate(headers):
        if needle == header:
            return idx
    for idx, header in enumerate(headers):
        if needle in header:
            return idx
    return None


def load_accident_frames(xlsx_path: Path) -> dict[tuple[str, str], int | None]:
    """Load accident frame indices from the XLSX annotations.

    Args:
        xlsx_path: Path to the DADA annotations spreadsheet.

    Returns:
        Mapping from (type_id, video_id) to accident frame index or None.
    """
    if not xlsx_path.exists():
        raise FileNotFoundError(f"XLSX not found: {xlsx_path}")
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    if "Sheet1" not in workbook.sheetnames:
        raise ValueError("Expected Sheet1 in XLSX annotations.")
    sheet = workbook["Sheet1"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("XLSX annotations are empty.")
    headers = [_normalize_header(value) for value in rows[0]]
    accident_idx = _find_header_index(headers, "accident frame")
    type_idx = _find_header_index(headers, "type")
    video_idx = _find_header_index(headers, "video")
    if accident_idx is None or type_idx is None or video_idx is None:
        raise ValueError("Required headers (type, video, accident frame) not found.")

    mapping: dict[tuple[str, str], int | None] = {}
    for row in rows[1:]:
        type_id = _normalize_type_id(row[type_idx])
        video_id = _normalize_video_id(row[video_idx])
        if not type_id or not video_id:
            continue
        accident_raw = row[accident_idx]
        accident_frame = None
        if accident_raw is not None:
            if isinstance(accident_raw, float) and accident_raw.is_integer():
                accident_raw = int(accident_raw)
            try:
                accident_frame = int(accident_raw)
            except (TypeError, ValueError):
                accident_frame = None
        if accident_frame is not None and accident_frame <= 0:
            accident_frame = None
        mapping[(type_id, video_id)] = accident_frame
    return mapping


def compute_max_frames(target_fps: float, max_seconds: float) -> int:
    """Compute maximum frames for a clip length budget."""
    if max_seconds <= 0:
        raise ValueError("max_seconds must be positive.")
    max_frames = int(round(target_fps * max_seconds))
    if max_frames <= 0:
        raise ValueError("max_frames must be positive.")
    return max_frames


def _get_gpu_name() -> str:
    if torch.cuda.is_available():
        return torch.cuda.get_device_name(0)
    return "cpu"


def resolve_batch_size(
    gpu_name: str,
    hardware_profiles: dict[str, int],
    default_batch: int,
) -> int:
    """Select batch size based on the current GPU name.

    Args:
        gpu_name: Detected GPU name string.
        hardware_profiles: Mapping of GPU name substrings to batch sizes.
        default_batch: Fallback batch size.

    Returns:
        Batch size to use for evaluation.
    """
    if default_batch <= 0:
        raise ValueError("default_batch must be positive.")
    if not hardware_profiles:
        return default_batch
    normalized = gpu_name.lower()
    for key, value in hardware_profiles.items():
        if key.lower() in normalized:
            return value
    return default_batch


def _attach_accident_frames(clips: list[ClipMeta], mapping: dict[tuple[str, str], int | None]) -> list[ClipMeta]:
    updated = []
    for clip in clips:
        accident_frame = mapping.get((clip.type_id, clip.video_id))
        updated.append(
            ClipMeta(
                type_id=clip.type_id,
                video_id=clip.video_id,
                clip_id=clip.clip_id,
                accident_frame=accident_frame,
            )
        )
    return updated


def _to_clip_key(clip: ClipMeta) -> str:
    return f"{clip.type_id}/{clip.video_id}"


def _accident_frame_to_seconds(accident_frame: int | None, source_fps: float) -> float | None:
    if accident_frame is None:
        return None
    return float(accident_frame) / source_fps


def _build_ground_truth(clip: ClipMeta, source_fps: float) -> dict[str, Any]:
    return {
        "accident_type": int(clip.type_id),
        "accident_frame_position_sec": _accident_frame_to_seconds(clip.accident_frame, source_fps),
    }


def _format_participants(participants: list[str]) -> str:
    """Convert participant identifiers into a readable phrase."""
    readable = [participant.replace("_", " ") for participant in participants]
    if readable == ["ego car", "car"]:
        readable = ["ego car", "another car"]
    if len(readable) == 2:
        return f"{readable[0]} and {readable[1]}"
    if len(readable) == 1:
        return readable[0]
    return ", ".join(readable)


def _format_class_summaries(mapping_path: Path, class_ids: list[int]) -> str:
    """Build a short summary list for selected class IDs."""
    raw = json.loads(mapping_path.read_text(encoding="utf-8"))
    lines = []
    for class_id in class_ids:
        entry = raw.get(str(class_id))
        if entry is None:
            raise ValueError(f"Class id {class_id} not found in mapping.")
        participants = entry.get("participants", [])
        interaction = entry.get("interaction", "unknown")
        participants_text = _format_participants(participants)
        interaction_text = interaction.replace("_", " ")
        lines.append(f"{class_id} -- participants -- {participants_text}, interaction -- {interaction_text}")
    return "\n".join(lines)


def _build_eval_prompt(schema_model: type[BaseModel], prompt_path: Path, mapping_path: Path) -> str:
    """Render the evaluation prompt template with schema and class summaries."""
    schema_json = json.dumps(schema_model.model_json_schema(), ensure_ascii=False)
    template_text = load_prompt_template(prompt_path)
    class_summaries = _format_class_summaries(mapping_path, POPULAR_CLASS_IDS)
    return render_prompt(
        template_text,
        {
            "json_schema": schema_json,
            "class_summaries": class_summaries,
        },
    )


def _batch_indices(total: int, batch_size: int) -> list[list[int]]:
    if batch_size <= 0:
        raise ValueError("batch_size must be positive.")
    indices = list(range(total))
    return [indices[i : i + batch_size] for i in range(0, total, batch_size)]


def _prepare_batch(
    clips: list[ClipMeta],
    indices: list[int],
    dataset_root: Path,
) -> tuple[list[ClipMeta], list[str]]:
    batch_clips = [clips[idx] for idx in indices]
    batch_video_paths = []
    existing_clips: list[ClipMeta] = []
    for clip in batch_clips:
        clip_dir = dataset_root / clip.type_id / clip.video_id
        video_path = clip_dir / "video.mp4"
        if not video_path.exists():
            logger.warning("Video not found; skipping clip {}: {}", _to_clip_key(clip), video_path)
            continue
        existing_clips.append(clip)
        batch_video_paths.append(str(video_path))
    return existing_clips, batch_video_paths


def _load_existing_results(output_path: Path) -> dict[str, dict[str, Any]]:
    if not output_path.exists():
        return {}
    try:
        raw = json.loads(output_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        logger.warning("Existing results at {} are invalid JSON; ignoring. ({})", output_path, exc)
        return {}
    if not isinstance(raw, dict):
        logger.warning("Existing results at {} are not a dictionary; ignoring.", output_path)
        return {}
    return raw


def _collect_completed_keys(results: dict[str, dict[str, Any]]) -> set[str]:
    return {clip_key for clip_key, record in results.items() if record.get("json_valid") is True}


def _filter_clips_by_keys(clips: list[ClipMeta], excluded_keys: set[str]) -> tuple[list[ClipMeta], int]:
    if not excluded_keys:
        return clips, 0
    remaining: list[ClipMeta] = []
    skipped = 0
    for clip in clips:
        if _to_clip_key(clip) in excluded_keys:
            skipped += 1
            continue
        remaining.append(clip)
    return remaining, skipped


def _merge_results(
    results: dict[str, dict[str, Any]],
    batch_clips: list[ClipMeta],
    outputs: list[InferenceOutcome],
    source_fps: float,
) -> None:
    for clip, outcome in zip(batch_clips, outputs):
        clip_key = _to_clip_key(clip)
        results[clip_key] = {
            "clip_id": clip.clip_id,
            "prediction": outcome.parsed,
            "json_valid": outcome.parsed is not None,
            "error": outcome.error,
            "attempts": outcome.attempts,
            "raw_text": outcome.raw_text,
            "ground_truth": _build_ground_truth(clip, source_fps),
        }


def _dump_results(output_path: Path, results: dict[str, dict[str, Any]]) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(results, ensure_ascii=False), encoding="utf-8")


def run_eval(
    config: EvalConfig,
    inferencer: QwenVideoInferencer | None = None,
    overwrite: bool = False,
) -> dict[str, dict[str, Any]]:
    """Run evaluation over a dataset split and return results.

    Args:
        config: Evaluation configuration.
        inferencer: Optional inferencer override for testing.

    Returns:
        Dictionary keyed by clip with prediction/metadata payloads.
    """
    load_dotenv(REPO_ROOT / ".env")
    clips = load_split(config.eval.split_path)
    accident_map = load_accident_frames(config.eval.xlsx_path)
    clips = _attach_accident_frames(clips, accident_map)

    existing_results = {} if overwrite else _load_existing_results(config.eval.output_path)
    completed_keys = _collect_completed_keys(existing_results)
    clips, skipped = _filter_clips_by_keys(clips, completed_keys)
    if skipped > 0:
        logger.info("Found {} already evaluated clips; proceeding with {} remaining.", skipped, len(clips))

    gpu_name = _get_gpu_name()
    batch_size = resolve_batch_size(
        gpu_name,
        config.hardware.gpu_batch_sizes,
        config.hardware.default_batch_size,
    )
    logger.info("Using batch size {} for GPU '{}'", batch_size, gpu_name)

    inferencer = inferencer or QwenVideoInferencer(
        model_id=config.model.id,
        device=config.model.device,
        dtype=config.model.dtype,
    )

    results: dict[str, dict[str, Any]] = dict(existing_results)
    prompt = _build_eval_prompt(DadaAccidentSchema, PROMPT_PATH, CLASS_MAPPING_PATH)
    max_frames = compute_max_frames(config.preprocess.target_fps, config.preprocess.max_seconds)
    batches = _batch_indices(len(clips), batch_size)
    for batch_idx, indices in enumerate(tqdm(batches, desc="Batches"), start=1):
        if config.eval.limit_batches is not None and batch_idx > config.eval.limit_batches:
            logger.info("Stopping after {} batches due to limit_batches.", config.eval.limit_batches)
            break
        batch_clips, batch_video_paths = _prepare_batch(clips, indices, config.eval.dataset_root)
        if not batch_clips:
            logger.warning("Skipping batch {} because no videos were available.", batch_idx)
            continue
        outputs = inferencer.infer_batch(
            batch_video_paths,
            prompt,
            DadaAccidentSchema,
            max_retries=config.eval.max_retries,
            max_new_tokens=config.model.max_new_tokens,
            sample_fps=config.preprocess.target_fps,
            max_frames=max_frames,
            max_pixels=config.preprocess.max_pixels,
            min_pixels=config.preprocess.min_pixels,
        )
        _merge_results(results, batch_clips, outputs, config.preprocess.source_fps)
        if batch_idx % config.eval.dump_every_batches == 0:
            _dump_results(config.eval.output_path, results)
    return results


def _build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Run DADA-2000 evaluation with batching.")
    parser.add_argument("--config", type=Path, default=REPO_ROOT / "task_2_evaluation_suite/eval_config.yaml", help="Path to eval YAML config.")
    parser.add_argument("--overwrite", action="store_true", help="Overwrite existing results.")
    return parser


def main() -> None:
    args = _build_arg_parser().parse_args()
    config = load_eval_config(args.config)
    results = run_eval(config, overwrite=args.overwrite)
    _dump_results(config.eval.output_path, results)
    logger.info("Wrote results to {}", config.eval.output_path)


if __name__ == "__main__":
    main()

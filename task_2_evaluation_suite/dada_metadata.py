from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from loguru import logger
from openpyxl import load_workbook

WEATHER_LABELS = {1: "sunny", 2: "rainy", 3: "snowy", 4: "foggy"}
LIGHT_LABELS = {1: "day", 2: "night"}
SCENE_LABELS = {1: "highway", 2: "tunnel", 3: "mountain", 4: "urban", 5: "rural"}
LINEAR_LABELS = {1: "arterials", 2: "curve", 3: "intersection", 4: "T-junction", 5: "ramp"}


@dataclass(frozen=True)
class DadaMetadata:
    """Parsed metadata for a single DADA-2000 clip."""

    type_id: str
    video_id: str
    accident_frame: int | None
    abnormal_start: int | None
    abnormal_end: int | None
    weather_id: int | None
    weather: str | None
    light_id: int | None
    light: str | None
    scene_id: int | None
    scene: str | None
    road_type_id: int | None
    road_type: str | None
    text: str | None
    cause: str | None
    measure: str | None


def normalize_header(header: str) -> str:
    """Normalize XLSX headers to a comparable form."""
    return " ".join(str(header).strip().lower().split())


def find_header_index(headers: list[str], needle: str) -> int | None:
    """Find header indices by exact or substring match."""
    for idx, header in enumerate(headers):
        if needle == header:
            return idx
    for idx, header in enumerate(headers):
        if needle in header:
            return idx
    return None


def normalize_int(value: Any) -> int | None:
    """Coerce numeric values to ints when possible."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def normalize_float(value: Any) -> float | None:
    """Coerce values to floats when possible."""
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def normalize_positive_int(value: Any) -> int | None:
    """Coerce values to positive integers."""
    parsed = normalize_int(value)
    if parsed is None or parsed <= 0:
        return None
    return parsed


def normalize_video_id(value: Any, target_width: int = 3) -> str:
    """Normalize a video ID by zero-padding numeric identifiers."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if text.isdigit():
        return text.zfill(target_width)
    return text


def normalize_type_id(value: Any) -> str | None:
    """Normalize type IDs to string form."""
    parsed = normalize_int(value)
    if parsed is None:
        return None
    return str(parsed)


def normalize_text(value: Any) -> str | None:
    """Normalize free-text fields."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def map_label(label_id: int | None, mapping: dict[int, str]) -> str | None:
    """Map categorical ids to labels."""
    if label_id is None:
        return None
    return mapping.get(label_id, "unknown")


@dataclass(frozen=True)
class DadaHeaderIndices:
    """Resolved header indices for DADA-2000 Sheet1."""

    type_idx: int
    video_idx: int
    accident_idx: int | None
    abnormal_start_idx: int | None
    abnormal_end_idx: int | None
    weather_idx: int | None
    light_idx: int | None
    scene_idx: int | None
    linear_idx: int | None
    texts_idx: int | None
    causes_idx: int | None
    measures_idx: int | None


def _require_header_index(headers: list[str], needle: str) -> int:
    index = find_header_index(headers, needle)
    if index is None:
        raise ValueError(f"Required header '{needle}' not found in XLSX.")
    return index


def _optional_header_index(headers: list[str], needles: list[str]) -> int | None:
    for needle in needles:
        index = find_header_index(headers, needle)
        if index is not None:
            return index
    return None


def _build_header_indices(headers: list[str]) -> DadaHeaderIndices:
    return DadaHeaderIndices(
        type_idx=_require_header_index(headers, "type"),
        video_idx=_require_header_index(headers, "video"),
        accident_idx=_optional_header_index(headers, ["accident frame"]),
        abnormal_start_idx=_optional_header_index(headers, ["abnormal start"]),
        abnormal_end_idx=_optional_header_index(headers, ["abnormal end"]),
        weather_idx=_optional_header_index(headers, ["weather"]),
        light_idx=_optional_header_index(headers, ["light"]),
        scene_idx=_optional_header_index(headers, ["scene", "scenes"]),
        linear_idx=_optional_header_index(headers, ["linear"]),
        texts_idx=_optional_header_index(headers, ["texts"]),
        causes_idx=_optional_header_index(headers, ["causes"]),
        measures_idx=_optional_header_index(headers, ["measures"]),
    )


def _load_sheet_rows(xlsx_path: Path) -> list[tuple[Any, ...]]:
    if not xlsx_path.exists():
        raise FileNotFoundError(f"XLSX not found: {xlsx_path}")
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    if "Sheet1" not in workbook.sheetnames:
        raise ValueError("Expected Sheet1 in XLSX annotations.")
    sheet = workbook["Sheet1"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("XLSX annotations are empty.")
    return rows


def _get_row_value(row: tuple[Any, ...], index: int | None) -> Any:
    if index is None or index >= len(row):
        return None
    return row[index]


def _build_metadata_entry(
    row: tuple[Any, ...],
    indices: DadaHeaderIndices,
    type_id: str,
    video_id: str,
) -> DadaMetadata:
    accident_frame = normalize_positive_int(_get_row_value(row, indices.accident_idx))
    abnormal_start = normalize_positive_int(_get_row_value(row, indices.abnormal_start_idx))
    abnormal_end = normalize_positive_int(_get_row_value(row, indices.abnormal_end_idx))
    weather_id = normalize_int(_get_row_value(row, indices.weather_idx))
    light_id = normalize_int(_get_row_value(row, indices.light_idx))
    scene_id = normalize_int(_get_row_value(row, indices.scene_idx))
    road_type_id = normalize_int(_get_row_value(row, indices.linear_idx))
    return DadaMetadata(
        type_id=type_id,
        video_id=video_id,
        accident_frame=accident_frame,
        abnormal_start=abnormal_start,
        abnormal_end=abnormal_end,
        weather_id=weather_id,
        weather=map_label(weather_id, WEATHER_LABELS),
        light_id=light_id,
        light=map_label(light_id, LIGHT_LABELS),
        scene_id=scene_id,
        scene=map_label(scene_id, SCENE_LABELS),
        road_type_id=road_type_id,
        road_type=map_label(road_type_id, LINEAR_LABELS),
        text=normalize_text(_get_row_value(row, indices.texts_idx)),
        cause=normalize_text(_get_row_value(row, indices.causes_idx)),
        measure=normalize_text(_get_row_value(row, indices.measures_idx)),
    )


def load_dada_metadata(xlsx_path: Path) -> dict[tuple[str, str], DadaMetadata]:
    """Load metadata entries from the DADA-2000 annotations spreadsheet.

    Args:
        xlsx_path: Path to the DADA annotations spreadsheet.

    Returns:
        Mapping from (type_id, video_id) to metadata.
    """
    rows = _load_sheet_rows(xlsx_path)
    headers = [normalize_header(value) for value in rows[0]]
    indices = _build_header_indices(headers)
    metadata: dict[tuple[str, str], DadaMetadata] = {}
    for row in rows[1:]:
        type_id = normalize_type_id(_get_row_value(row, indices.type_idx))
        video_id = normalize_video_id(_get_row_value(row, indices.video_idx))
        if not type_id or not video_id:
            continue
        key = (type_id, video_id)
        if key in metadata:
            logger.warning("Duplicate metadata row for type={} video={}; skipping duplicate.", type_id, video_id)
            continue
        metadata[key] = _build_metadata_entry(row, indices, type_id, video_id)
    return metadata

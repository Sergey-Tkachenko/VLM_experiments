from __future__ import annotations

import argparse
import json
import random
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

from dotenv import load_dotenv
from loguru import logger
from openpyxl import load_workbook

DEFAULT_XLSX_PATH = Path("/workspace/datasets/mm-au/dada_text_annotations.xlsx")
DEFAULT_SPLIT_PATHS = {
    "train": Path("/workspace/datasets/mm-au/train.json"),
    "val": Path("/workspace/datasets/mm-au/val.json"),
    "test": Path("/workspace/datasets/mm-au/test.json"),
}
DEFAULT_OUTPUT_DIR = Path("task_2_evaluation_suite")
DEFAULT_KEEP_TYPES = [
    "11",
    "43",
    "50",
    "10",
    "5",
    "6",
    "37",
    "48",
    "38",
    "8",
    "1",
    "57",
    "12",
    "49",
    "56",
    "14",
    "39",
    "42",
    "9",
]
DEFAULT_RATIOS = {"train": 3.0, "val": 1.0, "test": 3.0}
DEFAULT_SEED = 42
SPLIT_ORDER = ("train", "val", "test")


def _normalize_header(header: str) -> str:
    return " ".join(str(header).strip().lower().split())


def load_sheet1_rows(xlsx_path: Path) -> list[dict[str, Any]]:
    """Load Sheet1 annotations as list of row dictionaries."""
    if not xlsx_path.exists():
        raise FileNotFoundError(f"XLSX not found: {xlsx_path}")
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    if "Sheet1" not in workbook.sheetnames:
        raise ValueError("Expected Sheet1 in XLSX annotations.")
    sheet = workbook["Sheet1"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("XLSX annotations are empty.")
    headers = [_normalize_header(value) for value in rows[0]]
    parsed_rows = []
    for row in rows[1:]:
        row_data = {headers[idx]: value for idx, value in enumerate(row)}
        parsed_rows.append(row_data)
    return parsed_rows


def _normalize_video_id(value: Any, target_width: int = 3) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if text.isdigit():
        return text.zfill(target_width)
    return text


def _normalize_type_id(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    try:
        return str(int(value))
    except (TypeError, ValueError):
        return None


def normalize_type_and_video(row: dict[str, Any]) -> tuple[str, str] | None:
    """Normalize the type/video pair from an annotation row."""
    type_id = _normalize_type_id(row.get("type"))
    video_id = _normalize_video_id(row.get("video"))
    if not type_id or not video_id:
        return None
    return type_id, video_id


def _parse_split_item(item: list[Any]) -> tuple[str, str]:
    if not isinstance(item, list) or len(item) < 1:
        raise ValueError("Invalid split item structure.")
    pair = item[0]
    if not isinstance(pair, list) or len(pair) != 2:
        raise ValueError("Split item must contain [type_id, video_id].")
    type_id = str(pair[0]).strip()
    video_id = _normalize_video_id(pair[1])
    return type_id, video_id


def load_existing_splits(split_paths: dict[str, Path]) -> dict[str, list[tuple[str, str]]]:
    """Load existing splits and guard against duplicates across splits."""
    mapping: dict[tuple[str, str], str] = {}
    for split_name, path in split_paths.items():
        if not path.exists():
            raise FileNotFoundError(f"Split not found: {path}")
        data = json.loads(path.read_text(encoding="utf-8"))
        for item in data:
            pair = _parse_split_item(item)
            existing = mapping.get(pair)
            if existing and existing != split_name:
                logger.warning(
                    "Duplicate clip %s in splits %s/%s; assigning to test.",
                    pair,
                    existing,
                    split_name,
                )
                mapping[pair] = "test"
            else:
                mapping[pair] = split_name
    result = {name: [] for name in SPLIT_ORDER}
    for pair, split in mapping.items():
        result[split].append(pair)
    return result


def filter_pairs_by_type(
    pairs: Iterable[tuple[str, str]],
    keep_types: set[str],
) -> list[tuple[str, str]]:
    """Filter clip pairs to the allowed type IDs."""
    return [pair for pair in pairs if pair[0] in keep_types]


def build_all_pairs(
    rows: list[dict[str, Any]],
    keep_types: set[str],
) -> list[tuple[str, str]]:
    """Build all unique (type_id, video_id) pairs from Sheet1."""
    pairs = set()
    for row in rows:
        normalized = normalize_type_and_video(row)
        if normalized is None:
            continue
        type_id, video_id = normalized
        if type_id in keep_types:
            pairs.add((type_id, video_id))
    return sorted(pairs, key=lambda item: (int(item[0]), item[1]))


def allocate_counts_largest_remainder(total: int, ratios: dict[str, float]) -> dict[str, int]:
    """Allocate counts using the largest remainder method."""
    if total < 0:
        raise ValueError("Total must be non-negative.")
    ratio_sum = sum(ratios.values())
    if ratio_sum <= 0:
        raise ValueError("Ratios must sum to a positive value.")
    exacts = {name: total * (value / ratio_sum) for name, value in ratios.items()}
    floors = {name: int(value) for name, value in exacts.items()}
    remainder = total - sum(floors.values())
    if remainder == 0:
        return floors
    order_index = {name: idx for idx, name in enumerate(SPLIT_ORDER)}
    remainders = sorted(
        ((name, exacts[name] - floors[name]) for name in ratios),
        key=lambda item: (-item[1], order_index.get(item[0], len(order_index))),
    )
    for name, _ in remainders[:remainder]:
        floors[name] += 1
    return floors


def stratified_split_by_class(
    pairs: list[tuple[str, str]],
    ratios: dict[str, float],
    rng: random.Random,
) -> dict[str, list[tuple[str, str]]]:
    """Split pairs per class using per-class proportional allocation."""
    by_type: dict[str, list[tuple[str, str]]] = defaultdict(list)
    for pair in pairs:
        by_type[pair[0]].append(pair)
    result = {name: [] for name in SPLIT_ORDER}
    for type_id, items in sorted(by_type.items(), key=lambda item: int(item[0])):
        items.sort(key=lambda item: item[1])
        rng.shuffle(items)
        counts = allocate_counts_largest_remainder(len(items), ratios)
        cursor = 0
        for split_name in SPLIT_ORDER:
            take = counts.get(split_name, 0)
            if take:
                result[split_name].extend(items[cursor : cursor + take])
            cursor += take
    return result


def merge_existing_with_new(
    existing: dict[str, list[tuple[str, str]]],
    new: dict[str, list[tuple[str, str]]],
) -> dict[str, list[tuple[str, str]]]:
    """Merge existing split assignments with new ones."""
    merged = {name: [] for name in SPLIT_ORDER}
    seen: set[tuple[str, str]] = set()
    for split_name in SPLIT_ORDER:
        for pair in existing.get(split_name, []):
            if pair in seen:
                continue
            seen.add(pair)
            merged[split_name].append(pair)
        for pair in new.get(split_name, []):
            if pair in seen:
                continue
            seen.add(pair)
            merged[split_name].append(pair)
    return merged


def write_split_json(split: list[tuple[str, str]], output_path: Path) -> None:
    """Write a split JSON in the DADA-2000 format."""
    split_sorted = sorted(split, key=lambda item: (int(item[0]), item[1]))
    payload = [[[type_id, video_id], str(idx + 1).zfill(3)] for idx, (type_id, video_id) in enumerate(split_sorted)]
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=True, separators=(",", ":")), encoding="utf-8")


def _parse_keep_types(value: str) -> list[str]:
    items = [item.strip() for item in value.split(",")]
    return [item for item in items if item]


def _build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Build DADA-2000 splits with popular classes only.")
    parser.add_argument("--xlsx-path", type=Path, default=DEFAULT_XLSX_PATH, help="Path to XLSX annotations.")
    parser.add_argument("--train-split", type=Path, default=DEFAULT_SPLIT_PATHS["train"], help="Train split JSON.")
    parser.add_argument("--val-split", type=Path, default=DEFAULT_SPLIT_PATHS["val"], help="Val split JSON.")
    parser.add_argument("--test-split", type=Path, default=DEFAULT_SPLIT_PATHS["test"], help="Test split JSON.")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR, help="Output directory for splits.")
    parser.add_argument(
        "--keep-types",
        type=str,
        default=",".join(DEFAULT_KEEP_TYPES),
        help="Comma-separated list of type IDs to keep.",
    )
    parser.add_argument("--ratio-train", type=float, default=DEFAULT_RATIOS["train"], help="Train ratio weight.")
    parser.add_argument("--ratio-val", type=float, default=DEFAULT_RATIOS["val"], help="Val ratio weight.")
    parser.add_argument("--ratio-test", type=float, default=DEFAULT_RATIOS["test"], help="Test ratio weight.")
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED, help="Random seed for shuffling.")
    return parser


def main() -> None:
    repo_root = Path(__file__).resolve().parents[1]
    load_dotenv(repo_root / ".env")

    args = _build_arg_parser().parse_args()
    keep_types = set(_parse_keep_types(args.keep_types))
    ratios = {"train": args.ratio_train, "val": args.ratio_val, "test": args.ratio_test}

    rows = load_sheet1_rows(args.xlsx_path)
    all_pairs = build_all_pairs(rows, keep_types)

    existing = load_existing_splits(
        {
            "train": args.train_split,
            "val": args.val_split,
            "test": args.test_split,
        }
    )
    existing = {name: filter_pairs_by_type(pairs, keep_types) for name, pairs in existing.items()}

    existing_set = {pair for pairs in existing.values() for pair in pairs}
    remaining = [pair for pair in all_pairs if pair not in existing_set]

    rng = random.Random(args.seed)
    new_splits = stratified_split_by_class(remaining, ratios, rng)
    merged = merge_existing_with_new(existing, new_splits)

    output_dir = args.output_dir
    write_split_json(merged["train"], output_dir / "train.json")
    write_split_json(merged["val"], output_dir / "val.json")
    write_split_json(merged["test"], output_dir / "test.json")

    logger.info(
        "Wrote splits to {} (train={}, val={}, test={}).",
        output_dir,
        len(merged["train"]),
        len(merged["val"]),
        len(merged["test"]),
    )


if __name__ == "__main__":
    main()

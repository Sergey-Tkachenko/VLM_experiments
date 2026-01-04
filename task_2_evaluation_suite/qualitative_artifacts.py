from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
import tempfile
import textwrap
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from loguru import logger
from openpyxl import load_workbook


DEFAULT_DATASET_ROOT = Path("/workspace/datasets/mm-au/Origin/DADA2000/DADA2000")
DEFAULT_XLSX_PATH = Path("/workspace/datasets/mm-au/dada_text_annotations.xlsx")
DEFAULT_MAPPING_PATH = Path("task_2_evaluation_suite/type_id_to_parsed_description_mapping.json")
DEFAULT_OUTPUT_ROOT = Path("/workspace/VLM_experiments/artifacts")
DEFAULT_FPS = 30
DEFAULT_FONT_SIZE = 20
DEFAULT_WRAP_WIDTH = 100
WEATHER_LABELS = {1: "sunny", 2: "rainy", 3: "snowy", 4: "foggy"}
LIGHT_LABELS = {1: "day", 2: "night"}
SCENE_LABELS = {1: "highway", 2: "tunnel", 3: "mountain", 4: "urban", 5: "rural"}
LINEAR_LABELS = {1: "arterials", 2: "curve", 3: "intersection", 4: "T-junction", 5: "ramp"}


def _build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Generate qualitative artifacts for a DADA-2000 clip.")
    parser.add_argument("--clip", required=True, help="Clip folder path like 1/001.")
    parser.add_argument("--dataset-root", type=Path, default=DEFAULT_DATASET_ROOT, help="DADA-2000 root path.")
    parser.add_argument("--xlsx-path", type=Path, default=DEFAULT_XLSX_PATH, help="Path to XLSX annotations.")
    parser.add_argument("--mapping-path", type=Path, default=DEFAULT_MAPPING_PATH, help="Type mapping JSON path.")
    parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT_ROOT, help="Base output directory.")
    parser.add_argument("--fps", type=int, default=DEFAULT_FPS, help="Output FPS for MP4s.")
    return parser


def parse_clip_folder(clip_folder: str) -> tuple[int, str]:
    """Parse a clip folder string like '1/001' into (type_id, video_id)."""
    parts = Path(clip_folder).parts
    if len(parts) != 2:
        raise ValueError("clip must look like 'type/video', e.g. '1/001'.")
    type_id = int(parts[0])
    return type_id, parts[1]


def _normalize_header(header: str) -> str:
    return " ".join(str(header).strip().lower().split())


def load_annotations(xlsx_path: Path) -> list[dict[str, Any]]:
    """Load Sheet1 annotations as a list of row dictionaries."""
    if not xlsx_path.exists():
        raise FileNotFoundError(f"XLSX not found: {xlsx_path}")
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    if "Sheet1" not in workbook.sheetnames:
        raise ValueError("Expected Sheet1 in XLSX annotations.")
    sheet = workbook["Sheet1"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("XLSX annotations are empty.")
    headers = [_normalize_header(value) for value in rows[0]]
    parsed_rows = []
    for row in rows[1:]:
        row_data = {headers[idx]: value for idx, value in enumerate(row)}
        parsed_rows.append(row_data)
    return parsed_rows


def _normalize_video_id(value: Any, target_width: int) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if text.isdigit():
        return text.zfill(target_width)
    return text


def _normalize_type_id(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def select_annotation_row(rows: list[dict[str, Any]], type_id: int, video_id: str) -> dict[str, Any]:
    """Select the matching annotation row for the clip."""
    normalized_video = _normalize_video_id(video_id, len(video_id))
    matches = []
    for row in rows:
        row_type = _normalize_type_id(row.get("type"))
        row_video = _normalize_video_id(row.get("video"), len(video_id))
        if row_type == type_id and row_video == normalized_video:
            matches.append(row)
    if not matches:
        raise ValueError(f"No annotation row found for type={type_id}, video={video_id}.")
    if len(matches) > 1:
        logger.warning("Multiple annotation rows found; using the first match.")
    return matches[0]


def load_type_mapping(mapping_path: Path) -> dict[int, dict[str, Any]]:
    if not mapping_path.exists():
        raise FileNotFoundError(f"Mapping JSON not found: {mapping_path}")
    mapping_raw = json.loads(mapping_path.read_text(encoding="utf-8"))
    return {int(key): value for key, value in mapping_raw.items()}


def list_clip_frames(clip_dir: Path) -> list[Path]:
    frames_dir = clip_dir / "images"
    if not frames_dir.exists():
        raise FileNotFoundError(f"Frames directory not found: {frames_dir}")
    frame_paths = sorted(frames_dir.glob("*.png"))
    if not frame_paths:
        raise FileNotFoundError(f"No PNG frames found in {frames_dir}")
    return frame_paths


def slice_frames(frames: list[Path], start_frame: int, end_frame: int) -> list[Path]:
    start_idx = max(0, start_frame - 1)
    end_idx = min(len(frames), end_frame)
    if start_idx >= end_idx:
        raise ValueError("Abnormal window range is empty after bounds adjustment.")
    return frames[start_idx:end_idx]


def get_accident_frame(frames: list[Path], accident_frame: int | None) -> Path | None:
    if accident_frame is None:
        return None
    idx = accident_frame - 1
    if idx < 0 or idx >= len(frames):
        return None
    return frames[idx]


def build_overlay_text(row: dict[str, Any], type_mapping: dict[int, dict[str, Any]]) -> str:
    type_id = _normalize_type_id(row.get("type"))
    mapping = type_mapping.get(type_id, {})
    mapping_text = json.dumps(mapping, ensure_ascii=True)
    text = str(row.get("texts") or "").strip()
    cause = str(row.get("causes") or "").strip()
    measure = str(row.get("measures") or "").strip()
    weather_key = _normalize_header("weather(sunny,rainy,snowy,foggy)1-4")
    light_key = _normalize_header("light(day,night)1-2")
    scene_key = _normalize_header("scenes(highway,tunnel,mountain,urban,rural)1-5")
    linear_key = _normalize_header("linear(arterials,curve,intersection,T-junction,ramp) 1-5")

    weather_label = WEATHER_LABELS.get(_normalize_type_id(row.get(weather_key)), "unknown")
    light_label = LIGHT_LABELS.get(_normalize_type_id(row.get(light_key)), "unknown")
    scene_label = SCENE_LABELS.get(_normalize_type_id(row.get(scene_key)), "unknown")
    linear_label = LINEAR_LABELS.get(_normalize_type_id(row.get(linear_key)), "unknown")
    return (
        f"type: {type_id}\n"
        f"type_description: {mapping_text}\n"
        f"weather: {weather_label}\n"
        f"light: {light_label}\n"
        f"scene: {scene_label}\n"
        f"linear: {linear_label}\n"
        f"text: {text}\n"
        f"cause: {cause}\n"
        f"measure: {measure}"
    )


def _wrap_overlay_text(overlay_text: str, width: int = DEFAULT_WRAP_WIDTH) -> str:
    wrapped_lines = []
    for line in overlay_text.splitlines():
        wrapped_lines.extend(textwrap.wrap(line, width=width) or [""])
    return "\n".join(wrapped_lines)


def _find_font_path() -> Path | None:
    candidates = [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
        Path("/usr/share/fonts/truetype/freefont/FreeSans.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return None


def _escape_drawtext_path(path: Path) -> str:
    return str(path).replace("\\", "\\\\").replace(":", "\\:")


def _write_overlay_text(temp_dir: Path, overlay_text: str) -> Path:
    overlay_path = temp_dir / "overlay.txt"
    overlay_path.write_text(_wrap_overlay_text(overlay_text), encoding="utf-8")
    return overlay_path


def _build_drawtext_filter(text_path: Path, font_size: int) -> str:
    font_path = _find_font_path()
    textfile_value = _escape_drawtext_path(text_path)
    parts = [
        f"textfile={textfile_value}",
        f"fontsize={font_size}",
        "fontcolor=white",
        "box=1",
        "boxcolor=black@0.6",
        "boxborderw=8",
        "x=0",
        "y=0",
        "line_spacing=4",
    ]
    if font_path is not None:
        parts.insert(1, f"fontfile={_escape_drawtext_path(font_path)}")
    return "drawtext=" + ":".join(parts)


def _detect_image_codec(image_path: Path) -> str:
    with image_path.open("rb") as handle:
        signature = handle.read(8)
    if signature.startswith(b"\xff\xd8\xff"):
        return "mjpeg"
    if signature.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    return "png"


def write_mp4(frames: list[Path], output_path: Path, overlay_text: str, fps: int = DEFAULT_FPS) -> None:
    if shutil.which("ffmpeg") is None:
        raise RuntimeError("ffmpeg is required to write MP4 artifacts.")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    start_stem = frames[0].stem
    if not start_stem.isdigit():
        raise ValueError(f"Frame name is not numeric: {frames[0].name}")
    width = len(start_stem)
    frame_pattern = str(frames[0].parent / f"%0{width}d.png")
    input_codec = _detect_image_codec(frames[0])
    with tempfile.TemporaryDirectory() as temp_dir_str:
        overlay_path = _write_overlay_text(Path(temp_dir_str), overlay_text)
        drawtext_filter = _build_drawtext_filter(overlay_path, DEFAULT_FONT_SIZE)
        command = [
            "ffmpeg",
            "-y",
            "-r",
            str(fps),
            "-f",
            "image2",
            "-vcodec",
            input_codec,
            "-start_number",
            str(int(start_stem)),
            "-i",
            frame_pattern,
            "-frames:v",
            str(len(frames)),
            "-vf",
            drawtext_filter,
            "-c:v",
            "libx264",
            "-pix_fmt",
            "yuv420p",
            str(output_path),
        ]
        try:
            subprocess.run(command, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        except subprocess.CalledProcessError as exc:
            stderr_text = exc.stderr.decode("utf-8", errors="replace") if exc.stderr else ""
            logger.error("ffmpeg failed: {}", stderr_text.strip())
            raise RuntimeError("ffmpeg failed while writing MP4 artifact.") from exc


def write_accident_frame(image_path: Path, output_path: Path, overlay_text: str) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if shutil.which("ffmpeg") is None:
        raise RuntimeError("ffmpeg is required to write accident frame artifacts.")
    input_codec = _detect_image_codec(image_path)
    with tempfile.TemporaryDirectory() as temp_dir_str:
        overlay_path = _write_overlay_text(Path(temp_dir_str), overlay_text)
        drawtext_filter = _build_drawtext_filter(overlay_path, DEFAULT_FONT_SIZE)
        command = [
            "ffmpeg",
            "-y",
            "-f",
            "image2",
            "-vcodec",
            input_codec,
            "-i",
            str(image_path),
            "-vf",
            drawtext_filter,
            "-frames:v",
            "1",
            str(output_path),
        ]
        try:
            subprocess.run(command, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        except subprocess.CalledProcessError as exc:
            stderr_text = exc.stderr.decode("utf-8", errors="replace") if exc.stderr else ""
            logger.error("ffmpeg failed: {}", stderr_text.strip())
            raise RuntimeError("ffmpeg failed while writing accident frame artifact.") from exc


def _get_row_value(row: dict[str, Any], keys: list[str]) -> Any:
    for key in keys:
        normalized = _normalize_header(key)
        if normalized in row:
            return row.get(normalized)
    return None


def _parse_frame_value(row: dict[str, Any], keys: list[str]) -> int | None:
    value = _get_row_value(row, keys)
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def main() -> None:
    repo_root = Path(__file__).resolve().parents[1]
    load_dotenv(repo_root / ".env")

    args = _build_arg_parser().parse_args()
    type_id, video_id = parse_clip_folder(args.clip)
    clip_dir = args.dataset_root / str(type_id) / video_id

    rows = load_annotations(args.xlsx_path)
    row = select_annotation_row(rows, type_id, video_id)
    type_mapping = load_type_mapping(args.mapping_path)

    frames = list_clip_frames(clip_dir)
    overlay_text = build_overlay_text(row, type_mapping)

    abnormal_start = _parse_frame_value(row, ["abnormal start frame", "abnormal start"])
    abnormal_end = _parse_frame_value(row, ["abnormal end frame", "abnormal end"])
    if abnormal_start is None or abnormal_end is None:
        raise ValueError("Missing abnormal start/end in annotations.")

    abnormal_frames = slice_frames(frames, abnormal_start, abnormal_end)
    accident_frame_value = _parse_frame_value(row, ["accident frame"])
    accident_frame_path = get_accident_frame(frames, accident_frame_value)

    output_dir = args.output_root / f"{type_id}_{video_id}"
    full_clip_path = output_dir / "full_clip.mp4"
    abnormal_clip_path = output_dir / "abnormal_subclip.mp4"
    accident_frame_path_out = output_dir / "accident_frame.png"

    write_mp4(frames, full_clip_path, overlay_text, fps=args.fps)
    write_mp4(abnormal_frames, abnormal_clip_path, overlay_text, fps=args.fps)

    if accident_frame_path is None:
        logger.warning("No accident frame present for clip {}.", args.clip)
    else:
        write_accident_frame(accident_frame_path, accident_frame_path_out, overlay_text)

    logger.info("Artifacts written to {}", output_dir)


if __name__ == "__main__":
    sys.exit(main())
